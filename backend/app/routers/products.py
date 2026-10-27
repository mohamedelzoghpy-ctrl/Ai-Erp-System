"""Products and inventory (stock levels, adjustments, movement history)."""

from decimal import Decimal

from typing import List, Optional

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile, status
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from ..pagination import paginate
from ..audit import audit
from .. import exports
from ..deps import get_current_user, require_manager, user_role
from .. import models, schemas, stock
from ..database import get_db

router = APIRouter(prefix="/products", tags=["products & inventory"], dependencies=[Depends(get_current_user)])


def _product_or_404(db: Session, product_id: int) -> models.Product:
    product = db.get(models.Product, product_id)
    if product is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, f"Product {product_id} not found")
    return product


def _check_refs(db: Session, warehouse_id: Optional[int], supplier_id: Optional[int]):
    if warehouse_id is not None and db.get(models.Warehouse, warehouse_id) is None:
        raise HTTPException(422, f"Warehouse {warehouse_id} does not exist")
    if supplier_id is not None and db.get(models.Supplier, supplier_id) is None:
        raise HTTPException(422, f"Supplier {supplier_id} does not exist")


@router.post("", response_model=schemas.ProductOut, status_code=status.HTTP_201_CREATED)
def create_product(payload: schemas.ProductCreate, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    # Cost price and opening stock are manager-only (mirrors opening balances for parties).
    if user_role(user) not in ("admin", "manager"):
        if Decimal(str(payload.cost_price)) != 0:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "Only a manager or admin can set a cost price")
        if payload.quantity != 0:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "Only a manager or admin can set an opening stock balance")
    _check_refs(db, payload.warehouse_id, payload.supplier_id)
    data = payload.model_dump()
    opening = data.pop("quantity", 0)
    product = models.Product(**data)
    db.add(product)
    try:
        db.flush()
        if opening:
            home = stock.home_warehouse_id(db, product)
            if home is None:
                db.rollback()
                raise HTTPException(422, "Create a warehouse before adding opening stock")
            stock.move_stock(db, product, home, opening, "adjustment", "opening balance")
        audit(db, user, "create", "product", product, f"{product.sku} · {product.name}")
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status.HTTP_409_CONFLICT, f"SKU '{payload.sku}' already exists")
    db.refresh(product)
    return product


@router.get("", response_model=List[schemas.ProductOut])
def list_products(
    response: Response,
    search: Optional[str] = None,
    warehouse_id: Optional[int] = None,
    supplier_id: Optional[int] = None,
    low_stock: bool = False,
    is_active: Optional[bool] = None,
    stock_warehouse_id: Optional[int] = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db),
):
    """`stock_warehouse_id` adds `warehouse_quantity` (what is on the shelf in that warehouse) to each row."""
    query = db.query(models.Product)
    if search:
        like = f"%{search}%"
        query = query.filter(models.Product.name.ilike(like) | models.Product.sku.ilike(like))
    if warehouse_id is not None:
        query = query.filter(models.Product.warehouse_id == warehouse_id)
    if supplier_id is not None:
        query = query.filter(models.Product.supplier_id == supplier_id)
    if is_active is not None:
        query = query.filter(models.Product.is_active == is_active)
    if low_stock:
        query = query.filter(models.Product.quantity <= models.Product.reorder_level)
    rows = paginate(query, response, skip, limit, models.Product.id)
    if stock_warehouse_id is not None and rows:
        have = dict(
            db.query(models.ProductStock.product_id, models.ProductStock.quantity)
            .filter(models.ProductStock.warehouse_id == stock_warehouse_id, models.ProductStock.product_id.in_([p.id for p in rows]))
            .all()
        )
        for p in rows:
            p.warehouse_quantity = have.get(p.id, 0)
    return rows


@router.get("/export/xlsx")
def export_products_xlsx(db: Session = Depends(get_db)):
    products = db.query(models.Product).order_by(models.Product.id).all()
    rows = []
    total_qty = 0
    total_cost_value = 0
    total_sale_value = 0
    for p in products:
        wh = db.get(models.Warehouse, p.warehouse_id) if p.warehouse_id else None
        sup = db.get(models.Supplier, p.supplier_id) if p.supplier_id else None
        q = p.quantity or 0
        cost_value = q * float(p.cost_price)
        sale_value = q * float(p.sale_price)
        total_qty += q
        total_cost_value += cost_value
        total_sale_value += sale_value
        rows.append([
            ("text", p.sku),
            ("text", p.name),
            ("text", p.unit),
            ("money", p.cost_price),
            ("money", p.sale_price),
            ("int", p.quantity),
            ("int", p.reorder_level),
            ("money", cost_value),
            ("money", sale_value),
            ("text", wh.name if wh else ""),
            ("text", sup.name if sup else ""),
            ("text", "نعم" if p.is_active else "لا"),
        ])
    rows.append([
        ("text", "الإجمالي"),
        ("text", ""),
        ("text", ""),
        ("money", 0),
        ("money", 0),
        ("int", total_qty),
        ("int", 0),
        ("money", total_cost_value),
        ("money", total_sale_value),
        ("text", ""),
        ("text", ""),
        ("text", ""),
    ])
    content = exports.build_tables_xlsx(
        "قائمة المنتجات", "", [("المنتجات", ["الكود", "الاسم", "الوحدة", "سعر التكلفة", "سعر البيع", "الكمية", "حد إعادة الطلب", "قيمة التكلفة", "قيمة البيع", "المخزن", "المورد", "نشط"], rows)]
    )
    return Response(content=content, media_type=exports.XLSX_MIME, headers={"Content-Disposition": 'attachment; filename="products.xlsx"', "Cache-Control": "no-store"})


@router.post("/import/xlsx")
async def import_products_xlsx(file: UploadFile = File(...), db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    import io
    import openpyxl
    content = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(422, "Invalid Excel file")
    ws = wb.active
    # Find the header row (exported files have title/subtitle rows first).
    header_row_idx = 1
    for r in range(1, min(ws.max_row, 10) + 1):
        vals = [str(c.value).strip() if c.value is not None else "" for c in ws[r]]
        joined = " ".join(vals)
        if "الكود" in joined or "SKU" in joined or "sku" in joined.lower():
            header_row_idx = r
            break
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[header_row_idx]]

    def col(name: str):
        for i, h in enumerate(header):
            if name in h:
                return i
        return None

    sku_i = col("الكود")
    if sku_i is None:
        sku_i = col("SKU")
    if sku_i is None:
        sku_i = 0
    name_i = col("الاسم")
    if name_i is None:
        name_i = col("Name")
    unit_i = col("الوحدة")
    if unit_i is None:
        unit_i = col("Unit")
    cost_i = col("التكلفة")
    if cost_i is None:
        cost_i = col("Cost")
    sale_i = col("البيع")
    if sale_i is None:
        sale_i = col("Sale")
    reorder_i = col("إعادة الطلب")
    if reorder_i is None:
        reorder_i = col("Reorder")

    company_id = db.info["company_id"]
    updated = 0
    created = 0
    errors = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if not row or all(v is None for v in row):
            continue
        # Skip trailing totals row
        first = str(row[sku_i]).strip() if (sku_i is not None and sku_i < len(row) and row[sku_i] is not None) else ""
        if first in ("", "الإجمالي", "Total", "TOTAL"):
            continue
        sku = first
        if not sku:
            continue
        name = str(row[name_i]).strip() if (name_i is not None and name_i < len(row) and row[name_i] is not None) else sku
        product = db.query(models.Product).filter(models.Product.company_id == company_id, models.Product.sku == sku).one_or_none()
        is_new = product is None
        if is_new:
            product = models.Product(company_id=company_id, sku=sku, name=name, unit="قطعة", is_active=True)
            db.add(product)
            db.flush()
            created += 1
        else:
            updated += 1
        if name_i is not None and name_i < len(row) and row[name_i] is not None:
            product.name = str(row[name_i]).strip()
        if unit_i is not None and unit_i < len(row) and row[unit_i] is not None:
            product.unit = str(row[unit_i]).strip() or product.unit
        if cost_i is not None and cost_i < len(row) and row[cost_i] is not None:
            try:
                product.cost_price = Decimal(str(row[cost_i]))
            except Exception:
                errors.append(f"{sku}: invalid cost_price")
        if sale_i is not None and sale_i < len(row) and row[sale_i] is not None:
            try:
                product.sale_price = Decimal(str(row[sale_i]))
            except Exception:
                errors.append(f"{sku}: invalid sale_price")
        if reorder_i is not None and reorder_i < len(row) and row[reorder_i] is not None:
            try:
                product.reorder_level = int(row[reorder_i])
            except Exception:
                errors.append(f"{sku}: invalid reorder_level")
        # (counted above)
    db.commit()
    return {"updated": updated, "created": created, "errors": errors}


@router.get("/{product_id}", response_model=schemas.ProductOut)
def get_product(product_id: int, db: Session = Depends(get_db)):
    return _product_or_404(db, product_id)
    return _product_or_404(db, product_id)


@router.patch("/{product_id}", response_model=schemas.ProductOut)
def update_product(product_id: int, payload: schemas.ProductUpdate, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    product = _product_or_404(db, product_id)
    data = payload.model_dump(exclude_unset=True)
    # Cost price changes are manager-only (mirrors opening balances for parties).
    if "cost_price" in data and user_role(user) not in ("admin", "manager"):
        if Decimal(str(data["cost_price"])) != Decimal(str(product.cost_price)):
            raise HTTPException(status.HTTP_403_FORBIDDEN, "Only a manager or admin can change the cost price")
    _check_refs(db, data.get("warehouse_id"), data.get("supplier_id"))
    for key, value in data.items():
        setattr(product, key, value)
    audit(db, user, "update", "product", product, "، ".join(data) or None)
    db.commit()
    db.refresh(product)
    return product


@router.delete("/{product_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_product(
    product_id: int, db: Session = Depends(get_db), user: models.User = Depends(require_manager)
):
    product = _product_or_404(db, product_id)
    audit(db, user, "delete", "product", product.id, f"{product.sku} · {product.name}")
    db.delete(product)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            "Product has order or stock history; deactivate it instead (is_active=false)",
        )


@router.post("/{product_id}/adjust-stock", response_model=schemas.ProductOut)
def adjust_stock(product_id: int, payload: schemas.StockAdjustment, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """Manual correction (stock-take, damage, etc.). Never lets quantity go negative."""
    if payload.change == 0:
        raise HTTPException(422, "change must not be zero")
    product = (
        db.query(models.Product).filter(models.Product.id == product_id).with_for_update().one_or_none()
    )
    if product is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, f"Product {product_id} not found")
    warehouse_id = payload.warehouse_id if payload.warehouse_id is not None else stock.home_warehouse_id(db, product)
    if warehouse_id is None or db.get(models.Warehouse, warehouse_id) is None:
        raise HTTPException(422, "Choose an existing warehouse for this adjustment")
    stock.move_stock(db, product, warehouse_id, payload.change, "adjustment", payload.reference)
    audit(db, user, "adjust_stock", "product", product, f"{payload.change:+d} → {product.quantity}" + (f" ({payload.reference})" if payload.reference else ""))
    db.commit()
    db.refresh(product)
    return product


@router.get("/{product_id}/stock", response_model=List[schemas.ProductStockOut])
def product_stock(product_id: int, db: Session = Depends(get_db)):
    """Where this product's stock sits, one row per warehouse that holds some."""
    _product_or_404(db, product_id)
    rows = (
        db.query(models.ProductStock)
        .filter(models.ProductStock.product_id == product_id, models.ProductStock.quantity > 0)
        .order_by(models.ProductStock.warehouse_id)
        .all()
    )
    return [schemas.ProductStockOut(warehouse_id=r.warehouse_id, warehouse_name=r.warehouse.name, quantity=r.quantity) for r in rows]


@router.get("/{product_id}/movements", response_model=List[schemas.StockMovementOut])
def product_movements(product_id: int, db: Session = Depends(get_db)):
    _product_or_404(db, product_id)
    return (
        db.query(models.StockMovement)
        .filter(models.StockMovement.product_id == product_id)
        .order_by(models.StockMovement.id.desc())
        .all()
    )
